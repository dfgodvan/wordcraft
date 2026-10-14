# Synthetic DOCX for the table-of-contents update bug:
#   * a "Contents" paragraph (TOC Heading), then a TOC field laid out over paragraphs of its own
#     (begin + instruction + separate in a TOC 1 paragraph, a placeholder, end in a separate one);
#   * TOC 2 / TOC 3 styles with a hanging indent, a left tab stop for the number and a right
#     dot-leader stop; TOC 1 without stops of its own;
#   * short and long (wrapping) headings on three levels;
#   * a landscape last section (wider than the portrait ones).
#   make_toc_repro.py out.docx [--portrait-only]   (needs python-docx)
import sys
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import parse_xml
from docx.shared import Mm

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
d = Document()
styles = d.styles.element
for sid, name, ppr in [
    ('TOCHeading', 'TOC Heading', '<w:jc w:val="center"/>'),
    ('TOC1', 'toc 1', '<w:spacing w:after="100"/>'),
    ('TOC2', 'toc 2', '<w:tabs><w:tab w:val="left" w:pos="1540"/><w:tab w:val="right" w:leader="dot" w:pos="9344"/></w:tabs>'
                      '<w:spacing w:after="100"/><w:ind w:left="738" w:hanging="454"/>'),
    ('TOC3', 'toc 3', '<w:tabs><w:tab w:val="left" w:pos="1889"/><w:tab w:val="right" w:leader="dot" w:pos="9344"/></w:tabs>'
                      '<w:spacing w:after="100"/><w:ind w:left="1078" w:hanging="624"/>'),
]:
    styles.append(parse_xml(f'<w:style {W} w:type="paragraph" w:styleId="{sid}"><w:name w:val="{name}"/>'
                            f'<w:basedOn w:val="Normal"/><w:pPr>{ppr}</w:pPr></w:style>'))

for sec in d.sections:
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(30), Mm(15)

body = d.element.body
def raw(xml):
    body.insert(len(body) - 1, parse_xml(xml))

raw(f'<w:p {W}><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>Contents</w:t></w:r></w:p>')
raw(f'<w:p {W}><w:pPr><w:pStyle w:val="TOC1"/></w:pPr><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r></w:p>')
raw(f'<w:p {W}><w:r><w:t>No table of contents entries found.</w:t></w:r></w:p>')
raw(f'<w:p {W}><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')

long = ('A deliberately long section title that does not fit on one line of the table of contents '
        'and wraps onto a second and possibly a third line before its page number')
for i in range(1, 4):
    d.add_page_break()
    d.add_heading(f'{i} Section', 1)
    d.add_paragraph('Body text.')
    d.add_heading(f'{i}.1 Code', 2)
    d.add_paragraph('Body text.')
    d.add_heading(f'{i}.2 {long}', 2)
    d.add_paragraph('Body text.')
    d.add_heading(f'{i}.2.1 Parts', 3)
    d.add_paragraph('Body text.')
if '--portrait-only' not in sys.argv:
    # Landscape last section: the caret is here after Select All.
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    d.add_heading('4 Appendix', 1)
    d.add_paragraph('Body text.')
d.save(sys.argv[1])
